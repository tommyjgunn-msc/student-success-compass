"""ALCHE Model AU house style helpers for python-docx.

Matches ALCHE_MAU_CometMineral_Brief.pdf: Arial, navy banners, steel blue
headings, ALCHE red, pastel callout boxes, US Letter.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY      = "1A252F"
STEEL     = "1F618D"
GREEN     = "1E8449"
RED       = "922B21"
OLIVE     = "7D6608"
GREY      = "6B7280"
INK       = "1A252F"

T_GREEN   = "E9F7EF"
T_RED     = "FDEDEC"
T_CREAM   = "FEF9F0"
T_BLUE    = "EBF5FB"
T_PURPLE  = "F4ECF7"
T_GREY    = "F4F5F7"

FONT = "Arial"


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(INK)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.13
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    return doc


def _shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:color"), "auto"); el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def _cell_margins(cell, top=90, bottom=90, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{tag}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)


def _borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths, hexcolor) or None to hide."""
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{name}")
        spec = edges.get(name, "none")
        if spec == "none" or spec is None:
            e.set(qn("w:val"), "nil")
        else:
            size, color = spec
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size))
            e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tcPr.append(b)


def _para_border_left(p, size=18, color=STEEL, space=8):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), str(space)); e.set(qn("w:color"), color)
    b.append(e); pPr.append(b)


def _para_border_bottom(p, size=8, color=NAVY, space=6):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), str(space)); e.set(qn("w:color"), color)
    b.append(e); pPr.append(b)


def _run(p, text, size=10.5, bold=False, italic=False, color=INK, caps=False):
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.all_caps = caps
    r._element.rPr.rFonts.set(qn("w:cs"), FONT)
    return r


def _rich(p, text, size=10.5, color=INK, bold_all=False):
    """Renders **bold** and *italic* spans inline."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _run(p, part[2:-2], size=size, bold=True, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _run(p, part[1:-1], size=size, italic=True, bold=bold_all, color=color)
        else:
            _run(p, part, size=size, bold=bold_all, color=color)


# ---------- block elements ----------

def cover(doc, org, committee, sub, title, subtitle, meta):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; _shade(c, NAVY); _borders(c); _cell_margins(c, 320, 320, 200, 200)
    c.paragraphs[0].text = ""
    def line(txt, size, bold, color, space_after=2, caps=False):
        p = c.add_paragraph() if c.paragraphs[0].runs or c.paragraphs[0].text else c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        _run(p, txt, size=size, bold=bold, color=color, caps=caps)
    line(org, 11, True, "AED6F1", 2, True)
    line(committee, 8.5, False, "AAB7C4", 1, True)
    line(sub, 7.5, False, "8899A8", 10, True)
    line(title, 24, True, "FFFFFF", 4, True)
    line(subtitle, 11.5, True, "FFFFFF", 3)
    line(meta, 8.5, False, "8899A8", 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def statstrip(doc, items):
    """items: list of (big, small, colorhex)"""
    t = doc.add_table(rows=1, cols=len(items)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (big, small, col) in zip(t.rows[0].cells, items):
        _shade(cell, "FBFCFD"); _cell_margins(cell, 100, 100, 120, 120)
        _borders(cell, top=(6, "D5DBE1"), bottom=(6, "D5DBE1"), left=(6, "D5DBE1"), right=(6, "D5DBE1"))
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        _run(p, big, size=12.5, bold=True, color=col)
        p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        _run(p2, small, size=8.5, color=GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def part_head(doc, text, rule=True, new_page=False):
    p = doc.add_paragraph()
    if new_page:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
    else:
        p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6 if rule else 10)
    _run(p, text, size=16, bold=True, color="000000", caps=True)
    if rule:
        _para_border_bottom(p, size=8, color=NAVY, space=6)
    return p


def banner(doc, text, fill=NAVY, color="FFFFFF", new_page=False):
    if new_page:
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs: r.font.size = Pt(1)
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; _shade(c, fill); _borders(c); _cell_margins(c, 110, 110, 160, 160)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    _run(p, text, size=12, bold=True, color=color, caps=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def h2(doc, text, color=STEEL, size=12.5, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(4)
    _run(p, text, size=size, bold=True, color=color)
    return p


def h3(doc, text, color=STEEL, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    _run(p, text, size=size, bold=True, color=color, caps=True)
    return p


def body(doc, text, size=10.5, color=INK, space_after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    _rich(p, text, size=size, color=color)
    return p


def fact(doc, text):
    """The one-line takeaway under a heading. Steel bar, bold navy."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Inches(0.10)
    _para_border_left(p, size=18, color=STEEL, space=7)
    _rich(p, text, size=10.5, color=NAVY, bold_all=True)
    return p


def callout(doc, label, lines, fill=T_CREAM, bar=RED, label_color=None):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; _shade(c, fill); _cell_margins(c, 120, 120, 160, 150)
    _borders(c, left=(24, bar), top=(4, "E4E7EA"), bottom=(4, "E4E7EA"), right=(4, "E4E7EA"))
    first = True
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3 if i < len(lines) - 1 else 0)
        _rich(p, ln, size=10)
    if label:
        p0 = c.paragraphs[0].insert_paragraph_before()
        p0.paragraph_format.space_after = Pt(3)
        _run(p0, label, size=9, bold=True, color=label_color or bar, caps=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def clause_box(doc, number, label, text, fill, bar, label_color):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Inches(0.42); t.columns[1].width = Inches(6.08)
    cn, cb = t.rows[0].cells
    cn.width = Inches(0.42); cb.width = Inches(6.08)
    _shade(cn, bar); _borders(cn); _cell_margins(cn, 110, 110, 60, 60)
    pn = cn.paragraphs[0]; pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pn.paragraph_format.space_after = Pt(0)
    _run(pn, str(number), size=17, bold=True, color="FFFFFF")
    _shade(cb, fill); _cell_margins(cb, 110, 110, 150, 150)
    _borders(cb, top=(4, "E4E7EA"), bottom=(4, "E4E7EA"), right=(4, "E4E7EA"))
    p1 = cb.paragraphs[0]; p1.paragraph_format.space_after = Pt(3)
    _run(p1, label, size=9, bold=True, color=label_color, caps=True)
    p2 = cb.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    _rich(p2, text, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def grid(doc, headers, rows, widths=None, first_col_tint=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        if widths: c.width = Inches(widths[i])
        _shade(c, NAVY); _cell_margins(c, 80, 80, 120, 120)
        _borders(c, top=(6, NAVY), bottom=(6, NAVY), left=(6, NAVY), right=(6, NAVY))
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        _run(p, h, size=9.5, bold=True, color="FFFFFF")
    for r, row in enumerate(rows, start=1):
        for i, val in enumerate(row):
            c = t.rows[r].cells[i]
            if widths: c.width = Inches(widths[i])
            tint = T_BLUE if (i == 0 and first_col_tint) else "FFFFFF"
            _shade(c, tint); _cell_margins(c, 80, 80, 120, 120)
            _borders(c, top=(4, "D5DBE1"), bottom=(4, "D5DBE1"), left=(4, "D5DBE1"), right=(4, "D5DBE1"))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            _rich(p, val, size=9.5, color=NAVY if (i == 0 and first_col_tint) else INK,
                  bold_all=(i == 0 and first_col_tint))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.paragraph_format.space_after = Pt(3)
        _run(p, "•  ", size=size, color=STEEL, bold=True)
        _rich(p, it, size=size)


def spacer(doc, pts=8):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(pts)
    return p


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p
